from pptx import Presentation

prs = Presentation()

slides_data = [
    {
        "title": "DESIGN AND IMPLEMENTATION OF A REAL-TIME FAKE NEWS DETECTION SYSTEM",
        "content": "Ahmad Alhaji Ali — FCP/CSE/22/1005\nDepartment of Software Engineering | Federal University Dutse\nSupervisor: Mallam Umar Saminu Muhammad • Internal Oral Defence • 2025/2026"
    },
    {
        "title": "2. Introduction / Background",
        "content": "• Context: The digital age and social media have democratized information sharing, enabling the rapid spread of misinformation.\n• Importance: False content circulates faster than truth, threatening democratic processes and public health.\n• Current Situation: Traditional fact-checking (e.g., Dubawa) is accurate but inherently manual and reactive.\n• Technological Solution: Machine learning and NLP allow for automated text classification based on learned linguistic patterns.\n• Motivation: A critical need exists for an automated tool that flags false content in real-time, at the exact moment a reader encounters it."
    },
    {
        "title": "3. Problem Statement",
        "content": "• Problem 1: Existing standalone web applications require users to manually copy and paste text or URLs (workflow disruption).\n• Problem 2: Many automated tools act as a 'black-box,' providing a binary label without explaining the linguistic rationale.\n• Problem 3: Lack of transparency means users have little context or reason to trust the automated credibility scores.\n• Problem 4: High-accuracy models (e.g., Deep Learning) are rarely packaged into lightweight tools for ordinary readers.\n• Problem 5: A real-time, explainable solution is required to intercept misinformation directly within the browser ecosystem."
    },
    {
        "title": "4. Aim and Objectives",
        "content": "AIM:\nTo design and implement a real-time Google Chrome browser extension that detects fake news and provides explainable classifications using TF-IDF, SVM, and LIME.\n\nOBJECTIVES:\n• To review existing literature on fake news detection techniques and prepare the ISOT Fake News Dataset.\n• To design and train a TF-IDF + SVM classification model and integrate LIME for word-level explanations.\n• To develop a Flask REST API and a Chrome browser extension (Manifest V3) for real-time credibility assessments.\n• To evaluate the system using standard metrics and perform live testing on international and Nigerian news websites."
    },
    {
        "title": "5. Related Works — Literature Summary",
        "content": "• Karim et al. (2025): Used SVM + TF-IDF. Found SVM achieved highly competitive accuracy (99.6%+) but lacked an individual-prediction explainability component.\n• Hameed & Wasim (2025): Used LIME with BERT/RoBERTa. Found LIME provides justifications, enhancing user trust, but transformer models are computationally heavy.\n• Miao et al. (2023): Developed 'Factorial!' Chrome extension to help readers stay aware, but lacked LIME-based, word-level explanations.\n• Otor et al. (2023): Analyzed Nigerian fake-news dataset with CNN/RNN models. Found models can learn local contexts, but were not packaged as a deployed detection tool."
    },
    {
        "title": "6. Research Gap",
        "content": "• Existing approaches: High-accuracy classifiers are rarely deployed as tools an ordinary reader can use without manual copy-pasting.\n• Missing feature: Current browser tools often rely on human-curated publisher ratings rather than automated, per-article machine learning analysis.\n• Limitations: Explainability research (SHAP/LIME) is typically evaluated for technical consistency rather than packaged into an end-user-facing interface.\n• Specific context: Limited evidence combines a lightweight model (TF-IDF + SVM) with LIME in a deployed, real-time tool tested against Nigerian news.\n• Proposed Solution: Embedding a lightweight SVM with LIME explanations directly inside a real-time browsing experience."
    },
    {
        "title": "7. Methodology",
        "content": "REQUIREMENTS → DESIGN → DEVELOPMENT → TESTING → EVALUATION → DEPLOYMENT\n\n• Selected methodology: Design Science Research (DSR).\n• Data collection: Secondary Source (ISOT Fake News Dataset).\n• Justification: DSR is appropriate because the primary outcome is a working software artefact designed to address a practical problem, rather than just testing a theory.\n• Tools & Techniques: TF-IDF for numerical feature extraction, Random Undersampling for class balancing (42,834 articles), and Local Interpretable Model-agnostic Explanations (LIME) for transparency."
    },
    {
        "title": "8. Implementation / System Design",
        "content": "SYSTEM ARCHITECTURE:\nUser (Webpage) → Chrome Extension (Frontend) → Flask REST API (Business Logic) → SVM Pipeline (ML Processing) → JSON Response (Output)\n\nKey Interaction (Sequence Flow):\n1. The user clicks 'Analyze Current Page'.\n2. content.js extracts the active article text.\n3. Text is passed via POST request to the Flask API.\n4. Text is cleaned, vectorized, and classified by the Calibrated SVM.\n5. LIME generates local word weights.\n6. Results render on the popup UI."
    },
    {
        "title": "9. Implementation",
        "content": "• Frontend: HTML, CSS, JavaScript (Manifest V3)\n• Backend: Python (Flask REST API)\n• Database/Storage: Serialized ML models (.pkl via Joblib)\n• Framework / Tools: Scikit-learn, LIME, Pandas, NumPy\n\n(Speaker Note: Include screenshots of the Chrome Extension popup interface on this slide during the presentation.)"
    },
    {
        "title": "10. Results and Discussion",
        "content": "• Model Accuracy (Held-out Test Set): 99.78% (on 8,562 articles)\n• Response Target: Approx. < 2 seconds (Network dependent)\n• Real-World Test Cases: 17 external articles from 7 publishers\n• Explainability Integration: 100% of predictions accompanied by 10 LIME words\n\nDiscussion: While benchmark accuracy was exceptional (99.78%), real-world testing on contemporary publishers (like BBC and Yabaleftonline) showed varying confidence levels due to linguistic domain shifts, proving the necessity of the LIME explainability feature."
    },
    {
        "title": "11. Comparative Analysis with Existing Work",
        "content": "• Point-of-Reading Analysis: Manual (No) | Web Apps (No) | Proposed (Yes)\n• Automated Text Extraction: Manual (No) | Web Apps (No) | Proposed (Yes)\n• Word-Level Explainability: Manual (Yes, human) | Web Apps (No) | Proposed (Yes, LIME)\n• Automation: Manual (None) | Web Apps (Moderate) | Proposed (High)\n• Performance (Speed): Manual (Hours/Days) | Web Apps (Manual copy) | Proposed (Real-time < 2s)"
    },
    {
        "title": "12. Alignment of Findings with Research Objectives",
        "content": "• Objective 1: Review literature & prepare data. -> Achieved (ISOT dataset of 42,834 articles loaded, balanced via undersampling, and split).\n• Objective 2: Train TF-IDF + SVM with LIME. -> Achieved (Calibrated Linear SVM trained; LIME explainer integrated into pipeline).\n• Objective 3: Develop API & Extension. -> Achieved (Flask /predict endpoint and Manifest V3 extension built and integrated successfully).\n• Objective 4: Evaluate & perform live testing. -> Achieved (Evaluated metrics at 99.78% accuracy; tested on 17 live articles including Daily Trust and Yabaleftonline)."
    },
    {
        "title": "13. Summary",
        "content": "• The project investigated the rapid spread of misinformation and fake news in the digital ecosystem.\n• The major problem identified was the lack of real-time, point-of-reading tools that provide explainable AI justifications.\n• A software solution was designed using a client-server architecture and the Design Science Research methodology.\n• The system was implemented with a Python/Flask machine learning backend and a JavaScript Chrome Extension frontend.\n• Testing produced a 99.78% benchmark accuracy and successfully flagged unverified tabloid content (e.g., Yabaleftonline) in the wild.\n• The developed system addressed the 'black-box' AI problem by ensuring every prediction includes local LIME word weights."
    },
    {
        "title": "14. Conclusion",
        "content": "• The project successfully addressed the need for real-time fake news detection within the browser environment.\n• The proposed system was designed and implemented using highly efficient classical machine learning (TF-IDF + SVM) rather than resource-heavy deep learning.\n• Testing provided evidence that the system performs exceptionally well on benchmark data, while live observations confirmed its operational capability on external publishers.\n• The findings demonstrate that all four project objectives were fully achieved.\n• The system proves that providing users with an explainable, transparent AI interface enhances their ability to critically evaluate credibility signals."
    },
    {
        "title": "15. Recommendations",
        "content": "• Localisation: Train models on localized Nigerian news corpora to handle data distribution shifts and domain vocabulary.\n• Indigenous Language Support: Extend NLP techniques to support Hausa, Yoruba, Igbo, and Nigerian Pidgin.\n• Multimodal Detection: Integrate computer vision and audio analysis to evaluate the authenticity of deepfake images and videos.\n• Cross-Platform Expansion: Extend the extension to non-Chromium browsers (Firefox, Safari) and mobile platforms.\n• Fact-Checking Integration: Integrate live API calls to verified databases (like Dubawa) to cross-reference extracted claims.\n• Formal Benchmarking: Future researchers should conduct formal usability studies utilizing Human-Computer Interaction frameworks."
    },
    {
        "title": "16. THANK YOU",
        "content": "Questions & Answers\n\nAhmad Alhaji Ali • FCP/CSE/22/1005\nSupervisor: Mallam Umar Saminu Muhammad"
    }
]

for i, slide_dict in enumerate(slides_data):
    if i == 0:
        # Title Slide
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_dict["title"]
        slide.placeholders[1].text = slide_dict["content"]
    else:
        # Title and Content Slide
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_dict["title"]
        slide.placeholders[1].text = slide_dict["content"]

prs.save("Project_Defense_Presentation.pptx")
print("Presentation generated successfully: Project_Defense_Presentation.pptx")